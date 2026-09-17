# -*- coding: utf-8 -*-
"""オリアム株式会社 軽作業 50 件の元データから、同じ 50 勤務地の一般事務 50 行を組み立てる。

元データ (inputs/20260917-office-clerk/オリアム_軽作業50件_元データ.xlsx) の 66 列・列順・
シート名をそのまま使い、職種に依存する列だけを一般事務向けに書き換える。
勤務地 3 列とアクセスは元データの同じ行から引き継ぐので、駅名と住所がずれることはない。

会社の実態が分からない具体情報 (取扱商品、部署構成、使用ツール名) は書かない。
書けば虚偽の労働条件明示になる (職業安定法 65 条)。

    python3 gen.py
"""
from __future__ import annotations

import csv
import pathlib
import unicodedata

import openpyxl
from openpyxl.utils import get_column_letter

HERE = pathlib.Path(__file__).resolve().parent
REPO = HERE.parent.parent
SRC = REPO / "inputs/20260917-office-clerk/オリアム_軽作業50件_元データ.xlsx"
SHEET = "求人入力シート"
OUT_XLSX = HERE / "オリアム_一般事務50件_20260917.xlsx"
OUT_CSV = HERE / "オリアム_一般事務50件_20260917.csv"

TITLE_MAX_WIDTH = 30


def width(text: str) -> float:
    return sum(2 if unicodedata.east_asian_width(c) in "WFA" else 1 for c in text) / 2


# --- 書き換える列 -------------------------------------------------------------

JOB_CATEGORY = "一般事務"

# 職種名は元データの《未経験》検品・梱包作業/<駅名> と同じ並び (訴求語 → 職種 → 駅名) で、
# 職種の言い換えを 3 本回す。どれにも検索語「一般事務」を含める。
TITLE = [
    "《未経験OK》一般事務/{station}",
    "《未経験OK》一般事務スタッフ/{station}",
    "《未経験OK》データ入力・一般事務/{station}",
]

# キャッチコピーは 5 本を回す。給与額・年齢・性別には触れない。
# 書いている事実はすべて元データの列 (勤務時間・タグ・待遇・給与の補足・その他) にあるもの。
CATCH = [
    "未経験OK◎データ入力からはじめる一般事務♪",
    "残業なし・定時退社が基本◎はじめての事務職デビュー♪",
    "9:00〜18:00の固定時間◎生活リズムを崩さないオフィスワーク♪",
    "服装・ネイル自由◎自分らしく働ける一般事務スタッフ♪",
    "面接1回・履歴書任意◎未経験から事務職をはじめませんか♪",
]

# 仕事内容は 3 本を回す。駅名を差し込むので 50 行すべて別文面になる。
# 業務は一般事務の一般的な範囲 (データ入力、書類作成・ファイリング、電話・メール対応、
# 来客対応、備品管理) にとどめ、会社固有の情報は書かない。
BODY = [
"""✦・━━━━━━━━━━━━━━━・✦
＼{station}駅エリアで一般事務を募集中／
未経験からはじめるオフィスワーク◎
✦・━━━━━━━━━━━━━━━・✦

「事務の仕事に興味はあるけど経験がない…」
「パソコンは文字入力くらいしかできない…」

そんな方も、ひとつずつ覚えていける環境です。
最初は決まった形式への入力からスタートします。

＝＝＝＝＝＝＝＝＝＝＝＝＝＝＝
【✨この職場の魅力✨】

✅ 残業なし・定時退社が基本
9:00〜18:00の固定時間制なので、
仕事終わりの予定も立てやすいです。

✅ 未経験からスタート
分からないことは、その場で確認しながら
少しずつできる範囲を広げていけます。

✅ 服装・ネイル自由
自分らしいスタイルで働けます。

＝＝＝＝＝＝＝＝＝＝＝＝＝＝＝
【✧お仕事内容✧】

■ データ入力
決まった形式に、数字や文字を入力します。

■ 書類の作成・ファイリング
書類を作成し、種類ごとに整理して保管します。

■ 電話・メール対応
問い合わせを受け、担当者へ取り次ぎます。

■ 来客対応
お客様をご案内し、お茶出しなどを行います。

■ 備品の管理・発注
文房具などの在庫を確認し、必要な分を注文します。

＝＝＝＝＝＝＝＝＝＝＝＝＝＝＝
【✨働きやすいポイント✨】

✔ 9:00〜18:00の固定時間制（休憩1時間）
✔ 残業なし・定時退社が基本
✔ 社会保険完備・交通費支給
✔ 昇給年1回・賞与年2回

「人をサポートする仕事がしたい」
「腰を据えてオフィスで働きたい」
そんな方をお待ちしています！

【勤務地】
{address}（{station}駅から徒歩圏内）""",

"""✦・━━━━━━━━━━━━━━━・✦
＼{station}駅エリアで募集中／
9:00〜18:00・残業なしの一般事務◎
✦・━━━━━━━━━━━━━━━・✦

「接客や販売から、事務職に変わりたい」
「毎日同じ時間に帰れる仕事がいい」

一般事務は、9:00〜18:00の固定時間制。
残業なしで定時退社が基本なので、
生活のリズムを崩さずに働けます。

＝＝＝＝＝＝＝＝＝＝＝＝＝＝＝
【✧1日の流れ（一例）✧】

■ 9:00 出社
メールと郵便物を確認し、その日の予定を整理します。

■ 午前
データ入力や書類の作成を進めます。
電話がかかってきたら、担当者へ取り次ぎます。

■ 昼休憩（1時間）

■ 午後
書類のファイリング、備品の在庫確認・発注。
来客があればご案内します。

■ 18:00 退社
翌日の分を確認して、定時で退社します。

＝＝＝＝＝＝＝＝＝＝＝＝＝＝＝
【✨はじめての方へ✨】

✔ 経験・資格は必要ありません
✔ 最初はデータ入力など、決まった形式の作業から
✔ 分からないことは、その場で確認できます
✔ 服装・ネイル自由

＝＝＝＝＝＝＝＝＝＝＝＝＝＝＝
【✧待遇✧】

✅ 社会保険完備・交通費支給
✅ 昇給年1回・賞与年2回
✅ 資格取得支援制度あり

【勤務地】
{address}（{station}駅から徒歩圏内）""",

"""✦・━━━━━━━━━━━━━━━・✦
＼{station}駅エリアで募集中／
まずはデータ入力から。一般事務スタッフ◎
✦・━━━━━━━━━━━━━━━・✦

事務の仕事は、覚えることが多そうに見えて
実際は「決まった手順を正確に繰り返す」ことが中心です。

はじめはデータ入力から。
慣れてきたら、書類作成や電話対応へと
少しずつ担当を広げていきます。

＝＝＝＝＝＝＝＝＝＝＝＝＝＝＝
【✧担当する業務✧】

■ データ入力：決まった形式に数字や文字を入力する
■ 書類作成：ひな形をもとに書類を作る
■ ファイリング：書類を種類ごとに整理・保管する
■ 電話・メール対応：用件を聞いて担当者へ取り次ぐ
■ 来客対応：お客様のご案内、お茶出し
■ 備品管理：文房具などの在庫確認と発注

どの業務も、手順は先輩に確認しながら進められます。
自分ひとりで判断しなければならない場面は多くありません。

＝＝＝＝＝＝＝＝＝＝＝＝＝＝＝
【✨はじめやすいポイント✨】

✔ 経験・資格は必要ありません
✔ 9:00〜18:00の固定時間制（休憩1時間）
✔ 残業なし・定時退社が基本
✔ 服装・ネイル自由
✔ 面接は1回、履歴書は任意

＝＝＝＝＝＝＝＝＝＝＝＝＝＝＝
【✧待遇✧】

✅ 社会保険完備・交通費支給
✅ 昇給年1回・賞与年2回
✅ 健康診断・インフルエンザ予防接種

「コツコツ正確に進める仕事が好き」
そんな方に向いています。

【勤務地】
{address}（{station}駅から徒歩圏内）""",
]

# アピールポイントは全行共通。書いている事実は元データの列にあるものだけ。
APPEAL = """✅【残業なし・定時退社が基本】9:00〜18:00の固定時間制。仕事終わりの予定が立てやすく、プライベートの時間を確保できます。
✅【未経験からスタート】最初はデータ入力など決まった形式の作業から。分からないことはその場で確認しながら、少しずつ担当を広げていけます。
✅【服装・ネイル自由】自分らしいスタイルで働けます。
✅【昇給年1回・賞与年2回】社会保険完備、交通費支給、住宅手当・家族手当・皆勤手当あり。
✅【面接1回・履歴書任意】まずは興味がある程度で構いません。お気軽にご応募ください。"""

# 求める人材。年齢制限は元データが例外事由 3 号のイを明示しているので、
# legal.md の書式で維持する (期間の定めのない正社員・職務経験不問なので成立する)。
# 経験・資格を要件に入れると例外事由が成立しなくなるため、必須条件は置かない。
REQUIREMENTS = """【応募資格】
・未経験OK、職務経験不問
・学歴・資格不問

例外事由3号のイにより、長期キャリア形成を図るため、
39歳以下を募集しております。

★歓迎条件
・パソコンでの文字入力に慣れている方
・接客・販売など、人と接する仕事の経験がある方

▼こんな方におすすめ▼
・コツコツ正確に進める作業が好きな方
・残業なしで、生活のリズムを保って働きたい方
・人をサポートする仕事にやりがいを感じる方"""

# その他。元データの選考案内 (面接1回で内定) を活かし、末尾に変更の範囲を入れる。
# 変更の範囲は元データに情報がないため「変更なし」で置き、README で要確認に挙げる。
OTHER = """【選考について】
最短面接1回で内定です。
まずは興味がある程度で構いません。
お気軽にご応募ください。

＜選考の手順＞
【1】応募（履歴書は任意）
【2】面接（1回）
【3】内定

＝＝＝＝＝＝＝＝＝＝
＜労働条件の明示＞
■ 就業場所の変更の範囲
変更なし（本求人の勤務地のみ）
■ 従事すべき業務の変更の範囲
変更なし（本求人の業務のみ）"""

# タグ。元データの 3 つに、給与の補足「昇給年1回」に対応する「昇給・昇格あり」を足す。
TAGS = "残業なし、交通費支給、即日勤務OK、昇給・昇格あり"

# 自動アプローチ条件。元データの YAML 構造を保ち、キーワードだけ事務系に置き換える。
AUTO_APPROACH = (
    "preferred:\n"
    "  keywords:\n"
    "    professional: [一般事務, 事務補助, データ入力, 書類作成, ファイリング, 電話応対, 来客対応, 備品管理,\n"
    "      郵便物対応, スケジュール管理, 伝票処理, 経費精算, 庶務, OA事務, 受付事務, パソコン入力]\n"
)

# 引き継ぐ列のうち、倉庫作業に固有の記述だけを最小限直す。
# 「空調完備」は倉庫作業の訴求なので、オフィス勤務の一般事務からは外す。
BENEFITS_REMOVE_LINES = ("・空調完備",)

REWRITE_COLUMNS = (
    "職種名", "職業カテゴリー", "求人キャッチコピー",
    "募集要項（仕事内容）", "募集要項（アピールポイント）", "募集要項（求める人材）",
    "募集要項（その他）", "タグ", "自動アプローチ条件設定", "掲載画像",
)
ADJUST_COLUMNS = ("募集要項（待遇・福利厚生）",)


def station_of(title: str) -> str:
    """元データの職種名 `《未経験》検品・梱包作業/<駅名>` から駅名を取る。"""
    return title.rsplit("/", 1)[-1].strip()


def load_source() -> tuple[list[str], list[dict[str, str]]]:
    wb = openpyxl.load_workbook(SRC)
    ws = wb[SHEET]
    header = [str(c.value) for c in ws[1]]
    rows = []
    for cells in ws.iter_rows(min_row=2, values_only=True):
        if not any(v not in (None, "") for v in cells):
            continue
        rows.append({h: ("" if v is None else str(v)) for h, v in zip(header, cells)})
    return header, rows


def build_row(src: dict[str, str], index: int) -> dict[str, str]:
    station = station_of(src["職種名"])
    address = src["勤務地（都道府県・市区町村・町域）"]
    access = src["募集要項（アクセス）"]
    if station not in access:
        raise ValueError(f"行{index + 2}: 職種名の駅「{station}」がアクセス「{access}」に無い")

    row = dict(src)
    row["職種名"] = TITLE[index % len(TITLE)].format(station=station)
    row["職業カテゴリー"] = JOB_CATEGORY
    row["求人キャッチコピー"] = CATCH[index % len(CATCH)]
    row["募集要項（仕事内容）"] = BODY[index % len(BODY)].format(station=station, address=address)
    row["募集要項（アピールポイント）"] = APPEAL
    row["募集要項（求める人材）"] = REQUIREMENTS
    row["募集要項（その他）"] = OTHER
    row["タグ"] = TAGS
    row["自動アプローチ条件設定"] = AUTO_APPROACH
    row["掲載画像"] = ""

    benefits = [
        line for line in src["募集要項（待遇・福利厚生）"].split("\n")
        if line.strip() not in BENEFITS_REMOVE_LINES
    ]
    row["募集要項（待遇・福利厚生）"] = "\n".join(benefits)

    title_width = width(row["職種名"])
    if title_width > TITLE_MAX_WIDTH:
        raise ValueError(f"行{index + 2}: 職種名が全角 {title_width:.0f} 文字")
    return row


def write_xlsx(header: list[str], rows: list[dict[str, str]]) -> None:
    src_ws = openpyxl.load_workbook(SRC)[SHEET]
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = SHEET
    ws.append(header)
    for row in rows:
        ws.append([row[h] if row[h] != "" else None for h in header])
    for i in range(1, len(header) + 1):
        letter = get_column_letter(i)
        if letter in src_ws.column_dimensions:
            ws.column_dimensions[letter].width = src_ws.column_dimensions[letter].width
    wb.save(OUT_XLSX)


def write_csv(header: list[str], rows: list[dict[str, str]]) -> None:
    with OUT_CSV.open("w", encoding="utf-8-sig", newline="") as fh:
        writer = csv.writer(fh)
        writer.writerow(header)
        for row in rows:
            writer.writerow([row[h] for h in header])


def verify(header: list[str], src_rows: list[dict[str, str]], rows: list[dict[str, str]]) -> None:
    """書き換え対象外の列が 1 文字も変わっていないことを全行・全列で検算する。"""
    assert len(rows) == len(src_rows) == 50, len(rows)
    assert len(header) == 66, len(header)
    for i, (s, r) in enumerate(zip(src_rows, rows), start=2):
        for h in header:
            if h in REWRITE_COLUMNS or h in ADJUST_COLUMNS:
                continue
            assert s[h] == r[h], f"行{i} 列「{h}」が変わっている"
    bodies = {r["募集要項（仕事内容）"] for r in rows}
    titles = {r["職種名"] for r in rows}
    assert len(bodies) == 50, "仕事内容に重複がある"
    assert len(titles) == 50, "職種名に重複がある"
    for c in CATCH:
        assert width(c) <= 30, f"キャッチコピーが長い: {c}"


def main() -> None:
    header, src_rows = load_source()
    rows = [build_row(src, i) for i, src in enumerate(src_rows)]
    verify(header, src_rows, rows)
    write_xlsx(header, rows)
    write_csv(header, rows)
    print(f"{OUT_XLSX.name}: {len(rows)} 行 × {len(header)} 列")
    print(f"{OUT_CSV.name}: {len(rows)} 行 × {len(header)} 列")


if __name__ == "__main__":
    main()
